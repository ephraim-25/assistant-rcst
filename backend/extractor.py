import io
import re
from typing import Tuple
from docx import Document
from pypdf import PdfReader

def extract_text_from_file(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    """
    Extrait le texte brut d'un fichier .docx, .pdf ou texte.
    Retourne (texte_nettoyé, extension).
    """
    ext = "." + filename.split(".")[-1].lower() if "." in filename else ""
    raw_text = ""

    if ext in [".docx", ".doc"]:
        doc = Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text:
                full_text.append(para.text)
        # Extraire également le contenu des tableaux Word
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    full_text.append(" | ".join(row_texts))
        raw_text = "\n".join(full_text)

    elif ext == ".pdf":
        reader = PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                pages_text.append(t)
        raw_text = "\n".join(pages_text)

    else:
        # Fallback pour .txt, .md, etc.
        try:
            raw_text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            raw_text = file_bytes.decode("latin-1", errors="replace")

    # Nettoyage uniforme
    cleaned = raw_text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "").strip()
    return cleaned, ext
